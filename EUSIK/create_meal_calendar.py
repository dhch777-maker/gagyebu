import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date, timedelta
import calendar

# === 설정 ===
BABY_BIRTH_DATE = date(2025, 10, 28)
MEAL_START_DATE = date(2026, 4, 19)  # 일요일
OUTPUT_FILE = "EUSIK/초기_이유식_캘린더_v3.xlsx"

# === 28일 식단 데이터 ===
meal_plan = {
    1:  {"밥": "쌀죽과 퀵오트밀(10→7배죽)\n(잡곡 50%, 미음 대신 죽으로)", "고기류": "", "반찬": ""},
    2:  {"밥": "쌀죽과 퀵오트밀\n(서서히 늘려서 1~2주에\n50% 잡곡밥도 가능)", "고기류": "", "반찬": ""},
    3:  {"밥": "쌀오트밀 죽\n+ 소고기 다짐\n(얇게 썰어줘도 됨)", "고기류": "", "반찬": ""},
    4:  {"밥": "쌀오트밀 죽\n(빠르면 1배죽→7배죽)", "고기류": "소고기", "반찬": ""},
    5:  {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "+청경채\n(스틱 형태/토핑/죽에 섞기)\n(또는 청경채+당근)"},
    6:  {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "청경채\n(또는 청경채+당근)"},
    7:  {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "청경채(또는 청경채+당근)\n+사과 익힌 것\n(퓌레/슬라이스)\n(사과+바나나도 가능)"},
    8:  {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "청경채(또는 청경채+당근)\n사과 익힌 것\n(사과+바나나)"},
    9:  {"밥": "쌀오트밀 죽", "고기류": "소고기\n+계란 완전히 익힌 것", "반찬": "청경채(또는 청경채+당근)\n사과 익힌 것\n(사과+바나나)"},
    10: {"밥": "쌀오트밀 죽", "고기류": "소고기\n(완자로 줘도 됨)", "반찬": "청경채(또는 청경채+당근)\n사과 익힌 것\n(사과+바나나)"},
    11: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "청경채(또는 청경채+당근)\n사과 익힌 것\n(사과+바나나)"},
    12: {"밥": "쌀오트밀 죽", "고기류": "소고기\n+계란", "반찬": "+시금치(또는 시금치+단호박)\n사과"},
    13: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "+양배추\n시금치 단호박"},
    14: {"밥": "쌀오트밀 죽\n+ 밀가루 첨가", "고기류": "소고기", "반찬": "청경채+당근\n(또는 시금치+단호박)\n사과"},
    15: {"밥": "쌀오트밀 죽", "고기류": "소고기, 계란", "반찬": "청경채+당근\n(또는 시금치+단호박)\n사과"},
    16: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "청경채+당근\n(또는 시금치+단호박)\n사과"},
    17: {"밥": "쌀오트밀 죽 + 밀가루", "고기류": "소고기", "반찬": "양배추, 청경채, 당근\n+토마토\n사과"},
    18: {"밥": "쌀오트밀 죽", "고기류": "소고기, 계란", "반찬": "양배추, 청경채, 당근\n토마토\n사과"},
    19: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "단호박, 시금치\n+땅콩 소스(1주일에 3회)\n토마토, 사과"},
    20: {"밥": "쌀오트밀 죽 + 밀가루", "고기류": "소고기", "반찬": "단호박, 시금치\n또는 청경채, 당근\n토마토, 사과"},
    21: {"밥": "쌀오트밀 죽", "고기류": "소고기, 계란", "반찬": "당근, 시금치\n땅콩 소스\n토마토, 사과"},
    22: {"밥": "쌀오트밀 죽", "고기류": "소고기\n+생선(1주일에 2회)", "반찬": "당근, 시금치\n토마토, 사과"},
    23: {"밥": "쌀오트밀 죽 + 밀가루", "고기류": "소고기, 계란", "반찬": "단호박, 시금치\n땅콩 소스\n토마토, 사과"},
    24: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "당근, 시금치\n토마토, 사과"},
    25: {"밥": "쌀오트밀 죽", "고기류": "소고기, 생선", "반찬": "단호박, 시금치\n+또는 아보카도+딸기"},
    26: {"밥": "쌀오트밀 죽 + 밀가루", "고기류": "소고기, 계란", "반찬": "당근, 시금치\n땅콩 소스\n또는 아보카도, 딸기"},
    27: {"밥": "쌀오트밀 죽", "고기류": "+돼지고기 또는 닭고기", "반찬": "당근, 시금치\n또는 아보카도, 딸기"},
    28: {"밥": "쌀오트밀 죽", "고기류": "소고기", "반찬": "당근, 시금치\n땅콩 소스\n또는 아보카도, 딸기"},
}

# === 스타일 ===
HEADER_FILL = PatternFill(start_color="8B4513", end_color="8B4513", fill_type="solid")
HEADER_FONT = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
DAY_NUM_FONT = Font(name="맑은 고딕", size=10, bold=True, color="333333")
DPLUS_FONT = Font(name="맑은 고딕", size=9, bold=True, color="C0392B")
LABEL_FONT = Font(name="맑은 고딕", size=8, bold=True, color="8B4513")
CONTENT_FONT = Font(name="맑은 고딕", size=8, color="333333")
TITLE_FONT = Font(name="맑은 고딕", size=16, bold=True, color="8B4513")
INFO_FONT = Font(name="맑은 고딕", size=10, color="666666")
WEEKEND_FILL = PatternFill(start_color="FFF5F0", end_color="FFF5F0", fill_type="solid")
TODAY_FILL = PatternFill(start_color="FFEAA7", end_color="FFEAA7", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D5D5D5"),
    right=Side(style="thin", color="D5D5D5"),
    top=Side(style="thin", color="D5D5D5"),
    bottom=Side(style="thin", color="D5D5D5"),
)
THICK_BORDER = Border(
    left=Side(style="medium", color="8B4513"),
    right=Side(style="medium", color="8B4513"),
    top=Side(style="medium", color="8B4513"),
    bottom=Side(style="medium", color="8B4513"),
)

WEEK_COLORS = [
    PatternFill(start_color="FAF0E6", end_color="FAF0E6", fill_type="solid"),  # 1주차
    PatternFill(start_color="F5F0E8", end_color="F5F0E8", fill_type="solid"),  # 2주차
    PatternFill(start_color="F0EDE4", end_color="F0EDE4", fill_type="solid"),  # 3주차
    PatternFill(start_color="EBE8E0", end_color="EBE8E0", fill_type="solid"),  # 4주차
    PatternFill(start_color="E6E3DC", end_color="E6E3DC", fill_type="solid"),  # 5주차
]

DOW_KR = ["월", "화", "수", "목", "금", "토", "일"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "초기 이유식 캘린더"

# 인쇄 설정
ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.page_setup.orientation = "landscape"

# 열 너비 (7열 = 월~일)
for col in range(1, 8):
    ws.column_dimensions[get_column_letter(col)].width = 22

# === 상단 정보 영역 ===
ws.merge_cells("A1:G1")
cell = ws["A1"]
cell.value = "🍚 초기 이유식 28일 식단 캘린더"
cell.font = TITLE_FONT
cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

ws.merge_cells("A2:C2")
ws["A2"].value = f"👶 아기 출생일: {BABY_BIRTH_DATE.strftime('%Y-%m-%d')}"
ws["A2"].font = Info_font = Font(name="맑은 고딕", size=10, color="666666")
ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

ws.merge_cells("D2:G2")
ws["D2"].value = f"🥄 이유식 시작일: {MEAL_START_DATE.strftime('%Y-%m-%d')}  (D+{(MEAL_START_DATE - BABY_BIRTH_DATE).days})"
ws["D2"].font = Font(name="맑은 고딕", size=10, color="666666")
ws["D2"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 25

# 빈 행
ws.row_dimensions[3].height = 10

# === 캘린더 생성 ===
# 이유식 기간이 걸치는 월을 계산
end_date = MEAL_START_DATE + timedelta(days=27)

# 해당 기간의 모든 월 구하기
months = []
current = MEAL_START_DATE.replace(day=1)
while current <= end_date:
    months.append((current.year, current.month))
    if current.month == 12:
        current = current.replace(year=current.year + 1, month=1)
    else:
        current = current.replace(month=current.month + 1)

row_cursor = 4

for year, month in months:
    # 월 헤더
    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=7)
    month_cell = ws.cell(row=row_cursor, column=1)
    month_cell.value = f"{year}년 {month}월"
    month_cell.font = Font(name="맑은 고딕", size=13, bold=True, color="FFFFFF")
    month_cell.fill = HEADER_FILL
    month_cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, 8):
        ws.cell(row=row_cursor, column=c).fill = HEADER_FILL
        ws.cell(row=row_cursor, column=c).border = THICK_BORDER
    ws.row_dimensions[row_cursor].height = 30
    row_cursor += 1

    # 요일 헤더
    for i, dow in enumerate(DOW_KR):
        cell = ws.cell(row=row_cursor, column=i + 1)
        cell.value = dow
        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="A0522D", end_color="A0522D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row_cursor].height = 22
    row_cursor += 1

    # 달력 본체 (각 주 = 날짜 행 + 내용 행)
    cal = calendar.monthcalendar(year, month)

    DATE_ROW_FILL = PatternFill(start_color="F5E6D3", end_color="F5E6D3", fill_type="solid")
    DATE_ROW_WEEKEND = PatternFill(start_color="F0D5C0", end_color="F0D5C0", fill_type="solid")
    DATE_BORDER = Border(
        left=Side(style="thin", color="D5D5D5"),
        right=Side(style="thin", color="D5D5D5"),
        top=Side(style="thin", color="D5D5D5"),
        bottom=Side(style="hair", color="D5D5D5"),
    )
    CONTENT_BORDER = Border(
        left=Side(style="thin", color="D5D5D5"),
        right=Side(style="thin", color="D5D5D5"),
        top=Side(style="hair", color="D5D5D5"),
        bottom=Side(style="thin", color="D5D5D5"),
    )

    for week_idx, week in enumerate(cal):
        week_fill = WEEK_COLORS[week_idx % len(WEEK_COLORS)]
        date_row = row_cursor
        content_row = row_cursor + 1

        for col_idx, day in enumerate(week):
            d_cell = ws.cell(row=date_row, column=col_idx + 1)
            c_cell = ws.cell(row=content_row, column=col_idx + 1)

            d_cell.border = DATE_BORDER
            c_cell.border = CONTENT_BORDER
            d_cell.alignment = Alignment(horizontal="left", vertical="center")
            c_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            if day == 0:
                empty_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
                d_cell.fill = empty_fill
                c_cell.fill = empty_fill
                continue

            current_date = date(year, month, day)
            d_plus = (current_date - BABY_BIRTH_DATE).days
            meal_day = (current_date - MEAL_START_DATE).days + 1

            is_weekend = col_idx >= 5

            # 날짜 행: 배경색 구분
            if is_weekend:
                d_cell.fill = DATE_ROW_WEEKEND
                c_cell.fill = WEEKEND_FILL
            else:
                d_cell.fill = DATE_ROW_FILL
                c_cell.fill = week_fill

            if 1 <= meal_day <= 28:
                # 날짜 행
                d_cell.value = f"{day}일  D+{d_plus}  [{meal_day}일차]"
                d_cell.font = Font(name="맑은 고딕", size=9, bold=True, color="333333")

                # 내용 행
                meal = meal_plan[meal_day]
                lines = []
                if meal["밥"]:
                    lines.append(f"🍚 {meal['밥'].replace(chr(10), ' ')}")
                if meal["고기류"]:
                    lines.append(f"🥩 {meal['고기류'].replace(chr(10), ' ')}")
                if meal["반찬"]:
                    lines.append(f"🥬 {meal['반찬'].replace(chr(10), ' ')}")
                c_cell.value = "\n".join(lines)
                c_cell.font = Font(name="맑은 고딕", size=8, color="333333")
            else:
                # 이유식 기간 외 날짜
                d_cell.value = f"{day}일  D+{d_plus}"
                d_cell.font = Font(name="맑은 고딕", size=9, color="AAAAAA")
                c_cell.value = ""

        ws.row_dimensions[date_row].height = 22
        ws.row_dimensions[content_row].height = 80
        row_cursor += 2

    # 월 사이 간격
    ws.row_dimensions[row_cursor].height = 15
    row_cursor += 1

# === 하단 참고사항 ===
ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=7)
note_cell = ws.cell(row=row_cursor, column=1)
note_cell.value = "📌 참고사항"
note_cell.font = Font(name="맑은 고딕", size=11, bold=True, color="8B4513")
row_cursor += 1

notes = [
    "• 소고기는 매일, 닭고기는 어쩌다, 생선은 일주일에 2회 이하로 주는 것이 좋습니다.",
    "• 계란은 흰자 노른자 같이 줘도 됩니다.",
    "• 땅콩은 일주일에 3번 정도 주면 됩니다.",
    "• 한 번 첨가한 음식은 다음에 편하게 첨가해도 됩니다.",
    "• 2주가 지나면 양배추, 시금치, 당근 같은 상세 식재료를 기본 반찬으로 사용할 수 있습니다.",
    "• 미음으로 시작하지 말고 질감 있는 죽으로 시작하세요.",
    "• 이유식 초기에 잡곡을 50% 정도 첨가해서 먹여도 됩니다.",
]

for note in notes:
    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=7)
    cell = ws.cell(row=row_cursor, column=1)
    cell.value = note
    cell.font = Font(name="맑은 고딕", size=9, color="555555")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_cursor].height = 18
    row_cursor += 1

wb.save(OUTPUT_FILE)
print(f"✅ 캘린더 저장 완료: {OUTPUT_FILE}")
