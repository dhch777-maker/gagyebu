# -*- coding: utf-8 -*-
"""
원본 PPT 구조 그대로 유지하면서 컬러만 이마트 톤으로 전환.
컨셉: 검정(바탕) + 노랑(포인트) + 화이트(콘텐츠 카드/텍스트)

원본: 이마트에브리데이 사업제안서_260301.pptx (블루+네이비 톤, 흰 배경)
산출: 이마트에브리데이_사업제안서_v2_260301.pptx (검정 바탕, 노랑 포인트, 흰 카드)
"""
import sys
import shutil
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')

SRC = "이마트에브리데이 사업제안서_260301.pptx"
OUT = "이마트에브리데이_사업제안서_v2_260301.pptx"

# ============================================================
# 이마트 컨셉 컬러
# ============================================================
YELLOW = RGBColor(0xFF, 0xD2, 0x00)       # 이마트 옐로우
YELLOW_SOFT = RGBColor(0xFF, 0xEA, 0x80)   # 연노랑
BLACK = RGBColor(0x15, 0x15, 0x15)         # 바탕 검정
BLACK_DEEP = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xBB, 0xBB, 0xBB)
CARD_GRAY = RGBColor(0xEE, 0xEE, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ============================================================
# 원본 팔레트 → 신규 매핑
# 원본: 002B5B(네이비), 1F6BBA(블루), 74A9D8(연블루), F0F4F8(배경)
# ============================================================
# TEXT: 원본 네이비 텍스트는 제목(검정 배경 위) → 흰색
#       원본 블루는 강조 → 노랑
#       카드 안 회색/흰색 텍스트는 원래 밝은 카드 위였으므로 검정으로
TEXT_MAP = {
    "002B5B": WHITE,       # 네이비 제목 → 흰색 (검정 배경 위)
    "1F6BBA": YELLOW,      # 블루 강조 → 노랑
    "74A9D8": LIGHT_GRAY,  # 연블루 서브 → 라이트 그레이
    "666666": BLACK,       # 카드 내부 서브 텍스트 → 검정 (흰 카드 위)
    "FFFFFF": BLACK,       # 원본 네이비 블록 위 흰 글자 → 노랑 블록 위에선 검정
}

# FILL: 흰 카드 배경은 유지, 블루 강조는 노랑으로, 네이비 블록은 노랑으로
FILL_MAP = {
    "002B5B": YELLOW,      # 네이비 강조 블록 → 노랑
    "1F6BBA": YELLOW,      # 블루 강조 → 노랑
    "74A9D8": YELLOW_SOFT, # 연블루 → 연노랑
    "F0F4F8": WHITE,       # 연한 카드 배경 → 흰색
    "FFFFFF": WHITE,
}

LINE_MAP = {
    "002B5B": YELLOW,
    "1F6BBA": YELLOW,
    "74A9D8": LIGHT_GRAY,
    "DDDDDD": LIGHT_GRAY,
    "E5E5E5": LIGHT_GRAY,
}


def recolor_text(run):
    try:
        if run.font.color and run.font.color.type is not None:
            hex_col = str(run.font.color.rgb).upper()
            if hex_col in TEXT_MAP:
                run.font.color.rgb = TEXT_MAP[hex_col]
    except Exception:
        pass


def recolor_fill(shape):
    try:
        if shape.fill.type == 1:
            hex_col = str(shape.fill.fore_color.rgb).upper()
            if hex_col in FILL_MAP:
                shape.fill.fore_color.rgb = FILL_MAP[hex_col]
    except Exception:
        pass


def recolor_line(shape):
    try:
        if shape.line.fill.type == 1:
            hex_col = str(shape.line.color.rgb).upper()
            if hex_col in LINE_MAP:
                shape.line.color.rgb = LINE_MAP[hex_col]
    except Exception:
        pass


def walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            try:
                walk_shapes(shape.shapes)
            except Exception:
                pass
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    recolor_text(run)
        try:
            recolor_fill(shape)
        except Exception:
            pass
        try:
            recolor_line(shape)
        except Exception:
            pass
        if shape.has_table:
            tbl = shape.table
            for row in tbl.rows:
                for cell in row.cells:
                    try:
                        if cell.fill.type == 1:
                            hex_col = str(cell.fill.fore_color.rgb).upper()
                            if hex_col in FILL_MAP:
                                cell.fill.fore_color.rgb = FILL_MAP[hex_col]
                    except Exception:
                        pass
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            recolor_text(run)


def recolor_chart(shape):
    if not shape.has_chart:
        return
    chart = shape.chart
    try:
        for plot in chart.plots:
            for series in plot.series:
                try:
                    series.format.line.color.rgb = YELLOW
                    series.format.line.width = Pt(3.5)
                except Exception:
                    pass
                try:
                    series.marker.format.fill.solid()
                    series.marker.format.fill.fore_color.rgb = YELLOW
                    series.marker.format.line.color.rgb = WHITE
                    series.marker.size = 9
                except Exception:
                    pass
                try:
                    pts = list(series.points)
                    palette = [YELLOW, WHITE, YELLOW_SOFT, LIGHT_GRAY]
                    for i, pt in enumerate(pts):
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = palette[i % len(palette)]
                        try:
                            pt.format.line.color.rgb = BLACK
                        except Exception:
                            pass
                except Exception:
                    pass
                try:
                    plot.has_data_labels = True
                    dl = plot.data_labels
                    dl.font.bold = True
                    dl.font.color.rgb = WHITE
                    dl.font.size = Pt(11)
                except Exception:
                    pass
        # 축 글자 흰색화
        try:
            for ax in (chart.category_axis, chart.value_axis):
                ax.tick_labels.font.color.rgb = WHITE
                ax.tick_labels.font.size = Pt(10)
                ax.format.line.color.rgb = WHITE
        except Exception:
            pass
    except Exception as e:
        print(f"  chart recolor warn: {e}")


# ============================================================
# 배경을 검정으로 설정
# ============================================================
def set_black_background(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BLACK


# ============================================================
# 셰이프 유틸
# ============================================================
def add_rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    return shp


def move_to_back(shape):
    el = shape._element
    parent = el.getparent()
    parent.remove(el)
    parent.insert(2, el)


def add_text(slide, x, y, w, h, text, *, size=14, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = Pt(0); tf.margin_right = Pt(0)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


# ============================================================
# 이마트 프레임 추가 (상단 노랑 헤드라인 + 좌측 옐로우 바)
# ============================================================
def add_emart_frame(slide, idx, total, sw, sh):
    # 좌측 노랑 세로바
    left_bar = add_rect(slide, 0, 0, Inches(0.08), sh, YELLOW)
    # 상단 노랑 얇은 띠
    top_strip = add_rect(slide, 0, 0, sw, Inches(0.08), YELLOW)
    # 브랜드 식별
    brand = add_text(
        slide, Inches(0.3), Inches(0.25), Inches(5), Inches(0.3),
        "EMART EVERYDAY  ·  사업제안",
        size=9, bold=True, color=YELLOW, align=PP_ALIGN.LEFT,
    )
    # 페이지 번호
    page = add_text(
        slide, sw - Inches(1.5), Inches(0.25), Inches(1.2), Inches(0.3),
        f"{idx:02d} / {total:02d}",
        size=9, bold=True, color=YELLOW, align=PP_ALIGN.RIGHT,
    )


# ============================================================
# 표지 특별 스타일
# ============================================================
def style_cover(slide, sw, sh):
    # 상단 두꺼운 노랑 밴드 (인상 주기)
    band = add_rect(slide, 0, 0, sw, Inches(0.25), YELLOW)
    # 하단 얇은 노랑 띠
    bot = add_rect(slide, 0, sh - Inches(0.15), sw, Inches(0.15), YELLOW)
    # 중앙 포인트: 좌측 커다란 노랑 직사각형 블록
    block = add_rect(slide, 0, Inches(1.0), Inches(0.4), Inches(3.5), YELLOW)


def style_closing(slide, sw, sh):
    band_top = add_rect(slide, 0, 0, sw, Inches(0.15), YELLOW)
    band_bot = add_rect(slide, 0, sh - Inches(0.25), sw, Inches(0.25), YELLOW)


# ============================================================
# 메인
# ============================================================
def main():
    shutil.copy(SRC, OUT)
    prs = Presentation(OUT)
    sw = prs.slide_width
    sh = prs.slide_height
    total = len(prs.slides)

    for idx, slide in enumerate(prs.slides, start=1):
        # 1) 배경 검정화
        set_black_background(slide)

        # 2) 기존 셰이프 리컬러
        walk_shapes(slide.shapes)

        # 3) 차트 리컬러
        for shape in list(slide.shapes):
            recolor_chart(shape)

        # 4) 이마트 프레임 (좌측바 + 상단띠 + 페이지번호)
        add_emart_frame(slide, idx, total, sw, sh)

        # 5) 표지·마무리 특별 스타일
        if idx == 1:
            style_cover(slide, sw, sh)
        elif idx == total:
            style_closing(slide, sw, sh)

    prs.save(OUT)
    print(f"SAVED: {OUT}")
    print(f"SLIDES: {total}")


if __name__ == "__main__":
    main()
