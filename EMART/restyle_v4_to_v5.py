# -*- coding: utf-8 -*-
"""
v4 (사용자 수정본) → v5 증분 작업 전용 스크립트.

적용 범위 (그 외 슬라이드·레이아웃은 절대 건드리지 않음):
  1) 표지 슬라이드의 샛노랑(#FFD200) → 조금 더 찐한 연노랑(#FFE066)
     · 도형 fill, 텍스트 색
     · 우측 듀오톤 간판 이미지 재생성(같은 크기) 후 blob 교체
  2) 3p, 7p 프로세스 도식
     · OVAL 도형을 x 좌표 오름차순으로 정렬 → 노랑 전용 그라데이션(연→진)
     · RIGHT_ARROW 등 화살표 계열 도형 fill / line 을 검정(#151515)
"""
import os
import sys
import shutil
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from PIL import Image, ImageFilter

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "이마트에브리데이_사업제안서_v4_260301.pptx")
OUT = os.path.join(HERE, "이마트에브리데이_사업제안서_v5_260301.pptx")
SIGN_SRC = os.path.join(HERE, "이마트 간판.jpg")
BG_PATH = os.path.join(HERE, "_cover_bg_v5.png")

BLACK = RGBColor(0x15, 0x15, 0x15)
OLD_YELLOW_HEX = "FFD200"  # v4 표지의 샛노랑
NEW_COVER_YELLOW_RGB = (0xFF, 0xE0, 0x66)  # 조금 더 찐한 연노랑
NEW_COVER_YELLOW = RGBColor(*NEW_COVER_YELLOW_RGB)

# 노랑 전용 그라데이션 (오렌지/주황으로 안 넘어감, R=FF 고정)
GRAD_START = (0xFF, 0xF5, 0xB3)  # 매우 연한 노랑
GRAD_END = (0xFF, 0xC0, 0x00)    # 진한 노랑 (여전히 노랑)

ARROW_SHAPE_TYPES = {
    MSO_SHAPE.RIGHT_ARROW,
    MSO_SHAPE.LEFT_ARROW,
    MSO_SHAPE.UP_ARROW,
    MSO_SHAPE.DOWN_ARROW,
    MSO_SHAPE.LEFT_RIGHT_ARROW,
    MSO_SHAPE.PENTAGON,
    MSO_SHAPE.CHEVRON,
}


def _interp(c1, c2, t):
    return (
        int(round(c1[0] + (c2[0] - c1[0]) * t)),
        int(round(c1[1] + (c2[1] - c1[1]) * t)),
        int(round(c1[2] + (c2[2] - c1[2]) * t)),
    )


def yellow_only_gradient(n):
    return [
        RGBColor(*_interp(GRAD_START, GRAD_END, i / (n - 1) if n > 1 else 0))
        for i in range(n)
    ]


def _flatten(shapes):
    out = []
    for s in shapes:
        if s.shape_type == 6:
            try:
                out.extend(_flatten(s.shapes))
                continue
            except Exception:
                pass
        out.append(s)
    return out


def _hex(rgb):
    return str(rgb).upper() if rgb is not None else None


def tone_down_cover_fills_and_texts(slide):
    for s in _flatten(slide.shapes):
        try:
            if s.fill.type == 1 and _hex(s.fill.fore_color.rgb) == OLD_YELLOW_HEX:
                s.fill.fore_color.rgb = NEW_COVER_YELLOW
        except Exception:
            pass
        if hasattr(s, "has_text_frame") and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        if r.font.color and r.font.color.type is not None:
                            if _hex(r.font.color.rgb) == OLD_YELLOW_HEX:
                                r.font.color.rgb = NEW_COVER_YELLOW
                    except Exception:
                        pass


def _fit_crop(img, w, h):
    sw, sh = img.size
    tgt_ar = w / h
    src_ar = sw / sh
    if src_ar >= tgt_ar:
        scale = h / sh
        nw = int(round(sw * scale))
        img = img.resize((nw, h), Image.LANCZOS)
        left = (nw - w) // 2
        return img.crop((left, 0, left + w, h))
    else:
        scale = w / sw
        nh = int(round(sh * scale))
        img = img.resize((w, nh), Image.LANCZOS)
        top = (nh - h) // 2
        return img.crop((0, top, w, top + h))


def _duotone(img, dark, light):
    gray = img.convert("L")
    lut_r = [int(dark[0] + (light[0] - dark[0]) * i / 255) for i in range(256)]
    lut_g = [int(dark[1] + (light[1] - dark[1]) * i / 255) for i in range(256)]
    lut_b = [int(dark[2] + (light[2] - dark[2]) * i / 255) for i in range(256)]
    r = gray.point(lut_r)
    g = gray.point(lut_g)
    b = gray.point(lut_b)
    return Image.merge("RGB", (r, g, b))


def make_cover_background(src, out, w, h, light_rgb):
    img = Image.open(src).convert("RGB")
    img = _fit_crop(img, w, h)
    img = _duotone(img, (0x10, 0x10, 0x10), light_rgb)
    img = img.filter(ImageFilter.GaussianBlur(radius=max(2, int(h * 0.0025))))
    fade_w = int(w * 0.20)
    fade_row = Image.new("L", (w, 1), 255)
    for x in range(fade_w):
        fade_row.putpixel((x, 0), int(255 * (x / max(fade_w - 1, 1))))
    fade = fade_row.resize((w, h), Image.BILINEAR)
    bg = Image.new("RGB", (w, h), light_rgb)
    img = Image.composite(img, bg, fade)
    img.save(out, "PNG")


def retint_cover_picture(prs):
    slide = prs.slides[0]
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        return
    pic = pics[0]
    w, h = pic.image.size
    make_cover_background(SIGN_SRC, BG_PATH, w, h, NEW_COVER_YELLOW_RGB)
    blip = pic._element.blipFill.blip
    rId = getattr(blip, "rEmbed", None) or getattr(blip, "embed", None)
    img_part = slide.part.related_part(rId)
    with open(BG_PATH, "rb") as f:
        img_part._blob = f.read()


def apply_process_gradient(slide):
    ovals, arrows = [], []
    for s in _flatten(slide.shapes):
        try:
            ast = s.auto_shape_type
        except Exception:
            continue
        if ast == MSO_SHAPE.OVAL:
            ovals.append(s)
        elif ast in ARROW_SHAPE_TYPES:
            arrows.append(s)

    ovals.sort(key=lambda s: s.left if s.left is not None else 0)
    if len(ovals) >= 2:
        palette = yellow_only_gradient(len(ovals))
        for oval, col in zip(ovals, palette):
            oval.fill.solid()
            oval.fill.fore_color.rgb = col
            try:
                oval.line.color.rgb = BLACK
                oval.line.width = Pt(1.25)
            except Exception:
                pass

    for arr in arrows:
        try:
            arr.fill.solid()
            arr.fill.fore_color.rgb = BLACK
        except Exception:
            pass
        try:
            arr.line.color.rgb = BLACK
        except Exception:
            pass


def main():
    shutil.copy(SRC, OUT)
    prs = Presentation(OUT)

    # 1) 표지: fill·text 연노랑 + 듀오톤 이미지 재생성
    cover = prs.slides[0]
    tone_down_cover_fills_and_texts(cover)
    retint_cover_picture(prs)

    # 2) 3p, 7p: 그라데이션 + 화살표 검정
    apply_process_gradient(prs.slides[2])
    apply_process_gradient(prs.slides[6])

    prs.save(OUT)
    print(f"SAVED: {OUT}")
    print(f"SLIDES: {len(prs.slides)}")


if __name__ == "__main__":
    main()
